#!/usr/bin/env python3
"""
Extract all content from a .docx file with high accuracy.
Outputs both Markdown and JSON for downstream processing.

This script goes beyond basic python-docx by:
- Deep XML inspection to map images to their exact paragraph positions
- Heuristic caption association (figures and tables)
- Preserving style information for each paragraph
- Extracting table structure with merged cells
"""

import argparse
import json
import os
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent / "scripts"))
from common import fallback_output_dir  # noqa: E402


def extract_comments_data(doc):
    """
    Extract all comments from the comments.xml part.
    Returns a dict mapping comment id -> comment metadata.
    """
    comments_map = {}
    for rel in doc.part.rels.values():
        if "comments" in rel.reltype and "Extended" not in rel.reltype:
            comments_part = rel.target_part
            from docx.oxml import parse_xml
            root = parse_xml(comments_part.blob)
            for comment in root.findall(qn("w:comment")):
                cid = comment.get(qn("w:id"))
                author = comment.get(qn("w:author"))
                date = comment.get(qn("w:date"))
                initials = comment.get(qn("w:initials"))
                text = "".join(t.text or "" for t in comment.findall(".//" + qn("w:t")))
                comments_map[cid] = {
                    "id": cid,
                    "author": author,
                    "date": date,
                    "initials": initials,
                    "text": text,
                }
            break
    return comments_map


def extract_comment_ranges(doc):
    """
    Walk the document body in order to find which paragraphs each comment
    touches, and extract the text between commentRangeStart and commentRangeEnd.
    Returns a dict mapping comment id -> list of paragraph indices,
    and a dict mapping comment id -> selected text.
    """
    from docx.oxml import register_element_cls

    comment_paras = {}   # cid -> [para_idx, ...]
    comment_texts = {}   # cid -> selected text fragment

    body = doc.element.body
    para_idx = -1

    # Track open comments per paragraph for text extraction
    open_comments = set()

    for child in body:
        if child.tag == qn("w:p"):
            para_idx += 1
            # Reset open comments for each paragraph (range can span paras,
            # but text extraction is per-paragraph)
            open_comments = set()
            inside_range = False
            range_buffer = {}

            for sub in child:
                tag = sub.tag
                if tag == qn("w:commentRangeStart"):
                    cid = sub.get(qn("w:id"))
                    open_comments.add(cid)
                    comment_paras.setdefault(cid, []).append(para_idx)
                    range_buffer[cid] = []
                elif tag == qn("w:commentRangeEnd"):
                    cid = sub.get(qn("w:id"))
                    if cid in open_comments:
                        open_comments.discard(cid)
                        if cid in range_buffer:
                            frag = "".join(range_buffer[cid])
                            if frag:
                                if cid not in comment_texts:
                                    comment_texts[cid] = frag
                                else:
                                    comment_texts[cid] += frag
                            del range_buffer[cid]
                elif tag == qn("w:r"):
                    run_text = "".join(t.text or "" for t in sub.findall(qn("w:t")))
                    if run_text:
                        for cid in list(open_comments):
                            if cid in range_buffer:
                                range_buffer[cid].append(run_text)

            # Handle comments that end in a later paragraph
            for cid in open_comments:
                if cid in range_buffer:
                    frag = "".join(range_buffer[cid])
                    if frag:
                        if cid not in comment_texts:
                            comment_texts[cid] = frag
                        else:
                            comment_texts[cid] += frag

        elif child.tag == qn("w:tbl"):
            # Tables don't count as paragraphs but may be between paragraphs
            pass

    # Second pass: collect paragraph indices for comments that only have
    # commentReference (no range markers) or to ensure all referenced paras are captured.
    # We look at commentReference elements in each paragraph.
    para_idx = -1
    for child in body:
        if child.tag == qn("w:p"):
            para_idx += 1
            for sub in child.iter():
                if sub.tag == qn("w:commentReference"):
                    cid = sub.get(qn("w:id"))
                    if cid and cid not in comment_paras:
                        comment_paras[cid] = [para_idx]
                    elif cid and para_idx not in comment_paras[cid]:
                        comment_paras[cid].append(para_idx)

    return comment_paras, comment_texts


def extract_images(doc, docx_path):
    """
    Extract images with their paragraph positions and captions.
    Uses XML traversal to find exact paragraph index for each image.
    """
    images = []
    rels = doc.part.rels

    # Build a map of relationship IDs to image info
    image_rels = {}
    for rel in rels.values():
        if "image" in rel.reltype:
            image_rels[rel.rId] = {
                "rId": rel.rId,
                "target": rel.target_ref,
                "filename": os.path.basename(rel.target_ref),
            }

    # Walk all paragraphs to find images in their XML
    for para_idx, para in enumerate(doc.paragraphs):
        para_elem = para._element
        # Find all drawing elements in this paragraph
        drawings = para_elem.findall(".//" + qn("w:drawing"))
        blips = para_elem.findall(".//" + qn("a:blip"))

        for blip in blips:
            embed = blip.get(qn("r:embed"))
            if embed and embed in image_rels:
                img_info = image_rels[embed].copy()
                img_info["paragraph_index"] = para_idx
                img_info["paragraph_text"] = para.text
                images.append(img_info)

    # Associate captions with images
    # Strategy: look for figure/table captions near image paragraphs
    # Typical pattern: [body text referencing image] [image] [caption]
    for img in images:
        para_idx = img["paragraph_index"]
        caption = None
        caption_type = None

        # Look forward for caption (next 2 paragraphs)
        for offset in range(1, 3):
            check_idx = para_idx + offset
            if check_idx >= len(doc.paragraphs):
                break
            text = doc.paragraphs[check_idx].text.strip()
            if not text:
                continue
            # Match figure caption: "图X-Y ..." or "图X-Y"
            if re.match(r'^图\s*\d+([\.\-]\d+)?', text):
                caption = text
                caption_type = "figure"
                break
            # Match table caption: "表X-Y ..."
            elif re.match(r'^表\s*\d+([\.\-]\d+)?', text):
                caption = text
                caption_type = "table"
                break

        # If not found forward, look backward
        if not caption:
            for offset in range(1, 3):
                check_idx = para_idx - offset
                if check_idx < 0:
                    break
                text = doc.paragraphs[check_idx].text.strip()
                if not text:
                    continue
                if re.match(r'^图\s*\d+([\.\-]\d+)?', text):
                    caption = text
                    caption_type = "figure"
                    break
                elif re.match(r'^表\s*\d+([\.\-]\d+)?', text):
                    caption = text
                    caption_type = "table"
                    break

        img["caption"] = caption
        img["caption_type"] = caption_type

        # Try to extract figure/table number from caption
        if caption:
            m = re.search(r'[图表]\s*(\d+([\.\-]\d+)?)', caption)
            if m:
                img["number"] = m.group(1)

    return images


def _handle_merged_cells(table):
    """
    Parse Word Open XML cell properties (w:gridSpan, w:vMerge) and return
    a properly aligned 2D grid of cell texts.

    - w:gridSpan=N → content in first column, empty strings in N-1 spanned columns
    - w:vMerge restart → content carried forward to subsequent rows
    - w:vMerge continue → filled from the restart cell above

    Returns a list of lists, where each inner list has the same length (ncols).
    Non-merged tables return the same structure as the naive extraction.
    """
    ncols = 0
    for row in table.rows:
        pos = 0
        for cell in row.cells:
            tc_pr = cell._tc.find(qn("w:tcPr"))
            gs = 1
            if tc_pr is not None:
                gs_elem = tc_pr.find(qn("w:gridSpan"))
                if gs_elem is not None:
                    gs = int(gs_elem.get(qn("w:val"), "1"))
            pos += gs
        ncols = max(ncols, pos)

    grid = []
    vmerge_active = {}

    for row in table.rows:
        row_cells = [""] * ncols

        for col, entry in list(vmerge_active.items()):
            row_cells[col] = entry["content"]

        pos = 0
        for cell in row.cells:
            while pos < ncols and row_cells[pos]:
                pos += 1
            if pos >= ncols:
                break

            tc_pr = cell._tc.find(qn("w:tcPr"))
            text = cell.text.strip()
            gs = 1
            vmerge = None

            if tc_pr is not None:
                gs_elem = tc_pr.find(qn("w:gridSpan"))
                if gs_elem is not None:
                    gs = int(gs_elem.get(qn("w:val"), "1"))

                vm_elem = tc_pr.find(qn("w:vMerge"))
                if vm_elem is not None:
                    vm_val = vm_elem.get(qn("w:val"))
                    vmerge = "restart" if vm_val == "restart" else "continue"

            if vmerge == "restart":
                for i in range(gs):
                    vmerge_active[pos + i] = {"content": text, "gridSpan": gs}
                row_cells[pos] = text
            elif vmerge == "continue":
                if not row_cells[pos] and text:
                    row_cells[pos] = text
                    for i in range(gs):
                        vmerge_active[pos + i] = {"content": text, "gridSpan": gs}
            else:
                for i in range(gs):
                    vmerge_active.pop(pos + i, None)
                row_cells[pos] = text

            for i in range(1, gs):
                if pos + i < ncols:
                    row_cells[pos + i] = ""

            pos += gs

        grid.append(row_cells)

    return grid


def extract_tables(doc):
    """
    Extract tables with their positions, structure, and captions.
    """
    tables = []

    # We need to find where each table sits in the paragraph sequence
    # Tables are block-level elements in Word XML
    body = doc.element.body
    all_children = list(body)

    # Build a mapping: para_idx -> is it before/after a table?
    # More robust: iterate through body children in order
    para_idx = 0
    table_idx = 0

    for child in all_children:
        if child.tag == qn("w:p"):
            para_idx += 1
        elif child.tag == qn("w:tbl"):
            if table_idx < len(doc.tables):
                table = doc.tables[table_idx]
                table_data = {
                    "index": table_idx,
                    "paragraph_index": para_idx,
                    "rows": len(table.rows),
                    "columns": len(table.columns),
                    "cells": [],
                    "caption": None,
                }

                cells_grid = _handle_merged_cells(table)
                table_data["cells"] = cells_grid
                if cells_grid:
                    table_data["columns"] = len(cells_grid[0])

                # Find caption: look at paragraphs before the table
                # para_idx is the count of paragraphs BEFORE this table
                # so valid paragraph indices before table are 0 .. para_idx-1
                caption = None
                for offset in range(1, 5):
                    check_idx = para_idx - offset
                    if check_idx < 0:
                        break
                    text = doc.paragraphs[check_idx].text.strip()
                    if not text:
                        continue
                    if re.match(r'^表\s*\d+([\.\-]\d+)?', text):
                        caption = text
                        break

                # Also look after the table
                if not caption:
                    for offset in range(0, 5):
                        check_idx = para_idx + offset
                        if check_idx >= len(doc.paragraphs):
                            break
                        text = doc.paragraphs[check_idx].text.strip()
                        if not text:
                            continue
                        if re.match(r'^表\s*\d+([\.\-]\d+)?', text):
                            caption = text
                            break

                table_data["caption"] = caption
                if caption:
                    m = re.search(r'表\s*(\d+([\.\-]\d+)?)', caption)
                    if m:
                        table_data["number"] = m.group(1)

                tables.append(table_data)
                table_idx += 1

    return tables


def classify_heading_level(style_name):
    """
    Classify heading level from style name.
    Handles Chinese and English style names.
    """
    if not style_name:
        return 0

    style_lower = style_name.lower()

    # Direct heading levels
    if "heading 1" in style_lower or "标题 1" in style_name:
        return 1
    if "heading 2" in style_lower or "标题 2" in style_name:
        return 2
    if "heading 3" in style_lower or "标题 3" in style_name:
        return 3
    if "heading 4" in style_lower or "标题 4" in style_name:
        return 4
    if "heading 5" in style_lower or "标题 5" in style_name:
        return 5

    # 一级标题, 二级标题, etc.
    m = re.search(r'([一二三四五六七八九十]+)级标题', style_name)
    if m:
        chinese_nums = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5,
                        '六': 6, '七': 7, '八': 8, '九': 9, '十': 10}
        return chinese_nums.get(m.group(1), 0)

    return 0


def get_paragraph_format(para):
    """Extract formatting information from a paragraph."""
    fmt = {}
    pf = para.paragraph_format

    # Alignment
    if pf.alignment is not None:
        align_map = {0: "left", 1: "center", 2: "right", 3: "justify"}
        fmt["alignment"] = align_map.get(pf.alignment, str(pf.alignment))

    # Indentation
    if pf.first_line_indent:
        fmt["first_line_indent"] = str(pf.first_line_indent)
    if pf.left_indent:
        fmt["left_indent"] = str(pf.left_indent)

    # Spacing
    if pf.space_before:
        fmt["space_before"] = str(pf.space_before)
    if pf.space_after:
        fmt["space_after"] = str(pf.space_after)
    if pf.line_spacing:
        fmt["line_spacing"] = str(pf.line_spacing)

    return fmt


def get_run_formatting(run):
    """Extract formatting from a run."""
    fmt = {}
    font = run.font

    if font.name:
        fmt["font"] = font.name
    if font.size:
        fmt["size"] = str(font.size)
    if font.bold:
        fmt["bold"] = True
    if font.italic:
        fmt["italic"] = True
    if font.underline:
        fmt["underline"] = True
    if font.color and font.color.rgb:
        fmt["color"] = str(font.color.rgb)

    return fmt


def extract_paragraphs(doc, images, tables, comment_paras):
    """
    Extract all paragraphs with rich metadata.
    Mark positions of images and tables within the paragraph stream.
    """
    # Build lookup sets for quick checking
    image_para_indices = {img["paragraph_index"] for img in images}
    table_para_indices = {tbl["paragraph_index"] for tbl in tables}

    # Build reverse lookup: para_idx -> list of comment ids
    para_comments = {}
    for cid, para_indices in comment_paras.items():
        for pidx in para_indices:
            para_comments.setdefault(pidx, []).append(cid)

    paragraphs = []
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        style_name = para.style.name if para.style else None
        heading_level = classify_heading_level(style_name)

        p_data = {
            "index": idx,
            "text": text,
            "style": style_name,
            "heading_level": heading_level,
            "format": get_paragraph_format(para),
            "is_empty": not text,
            "has_image": idx in image_para_indices,
            "has_table": idx in table_para_indices,
            "comments": para_comments.get(idx, []),
        }

        # Extract run-level formatting
        runs_info = []
        for run in para.runs:
            run_fmt = get_run_formatting(run)
            if run_fmt or run.text:
                runs_info.append({
                    "text": run.text,
                    "format": run_fmt,
                })
        p_data["runs"] = runs_info

        paragraphs.append(p_data)

    return paragraphs


def detect_special_sections(paragraphs):
    """
    Detect special sections like abstract, keywords, references, etc.
    """
    sections = {
        "title": None,
        "abstract": None,
        "abstract_en": None,
        "keywords": None,
        "keywords_en": None,
        "references": [],
        "acknowledgements": None,
    }

    in_abstract = False
    in_abstract_en = False
    in_references = False
    in_ack = False

    abstract_lines = []
    abstract_en_lines = []
    reference_lines = []
    ack_lines = []

    for para in paragraphs:
        text = para["text"]
        style = para["style"] or ""
        heading = para["heading_level"]

        # Title: usually first heading 1
        if sections["title"] is None and heading == 1 and text:
            sections["title"] = text
            continue

        # Detect section transitions
        if heading >= 2:
            lower = text.lower()
            # Abstract
            if "摘要" in text and "abstract" not in lower:
                in_abstract = True
                in_abstract_en = False
                in_references = False
                in_ack = False
                continue
            elif "abstract" in lower and "摘要" not in text:
                in_abstract_en = True
                in_abstract = False
                in_references = False
                in_ack = False
                continue
            elif re.search(r'参考文[献献]', text) or "references" in lower:
                in_references = True
                in_abstract = False
                in_abstract_en = False
                in_ack = False
                continue
            elif "致谢" in text or "acknowledgement" in lower:
                in_ack = True
                in_abstract = False
                in_abstract_en = False
                in_references = False
                continue
            else:
                # Other section starts - reset all
                in_abstract = False
                in_abstract_en = False
                in_references = False
                in_ack = False
                continue

        # Collect content
        if in_abstract and text:
            abstract_lines.append(text)
        elif in_abstract_en and text:
            abstract_en_lines.append(text)
        elif in_references and text:
            reference_lines.append(text)
        elif in_ack and text:
            ack_lines.append(text)

        # Keywords detection
        if re.match(r'关键词[：:]\s*', text):
            sections["keywords"] = re.sub(r'关键词[：:]\s*', '', text)
        if re.match(r'[Kk]ey\s*[Ww]ords?[：:]\s*', text):
            sections["keywords_en"] = re.sub(r'[Kk]ey\s*[Ww]ords?[：:]\s*', '', text)

    sections["abstract"] = "\n".join(abstract_lines) if abstract_lines else None
    sections["abstract_en"] = "\n".join(abstract_en_lines) if abstract_en_lines else None
    sections["references"] = reference_lines
    sections["acknowledgements"] = "\n".join(ack_lines) if ack_lines else None

    return sections


def build_heading_tree(paragraphs):
    """Build a tree structure from heading paragraphs."""
    tree = []
    stack = []

    for para in paragraphs:
        level = para["heading_level"]
        if level == 0:
            continue

        node = {
            "level": level,
            "title": para["text"],
            "index": para["index"],
            "children": [],
        }

        # Find parent
        while stack and stack[-1]["level"] >= level:
            stack.pop()

        if stack:
            stack[-1]["children"].append(node)
        else:
            tree.append(node)

        stack.append(node)

    return tree


def generate_markdown(data):
    """Generate Markdown from extracted data."""
    lines = []
    paragraphs = data["paragraphs"]
    images = data["images"]
    tables = data["tables"]
    comments = data.get("comments", [])

    # Build lookup maps
    image_by_para = {img["paragraph_index"]: img for img in images}
    table_by_para = {tbl["paragraph_index"]: tbl for tbl in tables}
    comments_map = {c["id"]: c for c in comments}

    for para in paragraphs:
        idx = para["index"]
        text = para["text"]
        level = para["heading_level"]

        # Skip empty paragraphs
        if not text and not para["has_image"]:
            continue

        # Handle headings
        if level > 0:
            lines.append(f"{'#' * level} {text}")
            lines.append("")
            # Append comments for this heading
            for cid in para.get("comments", []):
                c = comments_map.get(cid)
                if c:
                    lines.append(f"> **批注** [{c['author']}]: {c['text']}")
                    if c.get("selected_text"):
                        lines.append(f"> *选中内容：{c['selected_text']}*")
                    lines.append("")
            continue

        # Handle images
        if para["has_image"] and idx in image_by_para:
            img = image_by_para[idx]
            if img.get("caption"):
                lines.append(f"**{img['caption']}**")
            lines.append(f"[图片: {img.get('filename', 'unknown')}]")
            lines.append("")
            # Don't output the empty paragraph text
            if text:
                lines.append(text)
                lines.append("")
            for cid in para.get("comments", []):
                c = comments_map.get(cid)
                if c:
                    lines.append(f"> **批注** [{c['author']}]: {c['text']}")
                    if c.get("selected_text"):
                        lines.append(f"> *选中内容：{c['selected_text']}*")
                    lines.append("")
            continue

        # Handle tables
        if para["has_table"] and idx in table_by_para:
            tbl = table_by_para[idx]
            if tbl.get("caption"):
                lines.append(f"**{tbl['caption']}**")

            # Generate markdown table
            cells = tbl["cells"]
            if cells:
                # Header row
                header = cells[0]
                lines.append("| " + " | ".join(header) + " |")
                lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                for row in cells[1:]:
                    lines.append("| " + " | ".join(row) + " |")
            lines.append("")
            for cid in para.get("comments", []):
                c = comments_map.get(cid)
                if c:
                    lines.append(f"> **批注** [{c['author']}]: {c['text']}")
                    if c.get("selected_text"):
                        lines.append(f"> *选中内容：{c['selected_text']}*")
                    lines.append("")
            continue

        # Regular paragraph
        if text:
            lines.append(text)
            lines.append("")
        for cid in para.get("comments", []):
            c = comments_map.get(cid)
            if c:
                lines.append(f"> **批注** [{c['author']}]: {c['text']}")
                if c.get("selected_text"):
                    lines.append(f"> *选中内容：{c['selected_text']}*")
                lines.append("")

    return "\n".join(lines)


def extract_all(docx_path):
    """Main extraction function."""
    doc = Document(docx_path)

    images = extract_images(doc, docx_path)
    tables = extract_tables(doc)
    comments_map = extract_comments_data(doc)
    comment_paras, comment_texts = extract_comment_ranges(doc)
    paragraphs = extract_paragraphs(doc, images, tables, comment_paras)
    sections = detect_special_sections(paragraphs)
    heading_tree = build_heading_tree(paragraphs)

    # Enrich comments with paragraph indices and selected text
    comments = []
    for cid, cdata in comments_map.items():
        comments.append({
            "id": cdata["id"],
            "author": cdata["author"],
            "date": cdata["date"],
            "initials": cdata["initials"],
            "text": cdata["text"],
            "paragraph_indices": comment_paras.get(cid, []),
            "selected_text": comment_texts.get(cid, None),
        })

    result = {
        "source_file": os.path.basename(docx_path),
        "statistics": {
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "image_count": len(images),
            "heading_count": sum(1 for p in paragraphs if p["heading_level"] > 0),
            "comment_count": len(comments),
        },
        "sections": sections,
        "heading_tree": heading_tree,
        "images": images,
        "tables": tables,
        "comments": comments,
        "paragraphs": paragraphs,
    }

    return result


def main():
    parser = argparse.ArgumentParser(description="Extract all content from a .docx file")
    parser.add_argument("input", help="Path to the .docx file")
    parser.add_argument("--output-dir", "-o", default=None, help="Output directory (default: auto-detect via two-level fallback)")
    parser.add_argument("--format", "-f", choices=["json", "markdown", "both"], default="both",
                        help="Output format")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"Error: File not found: {args.input}", file=sys.stderr)
        sys.exit(1)

    if not args.input.endswith(".docx"):
        print("Error: Input file must be a .docx file", file=sys.stderr)
        sys.exit(1)

    print(f"Extracting: {args.input}")
    data = extract_all(args.input)

    base_name = Path(args.input).stem
    if args.output_dir is not None:
        output_dir = Path(args.output_dir)
    else:
        # Two-level fallback（OUTPUT.md C-1）：基准是主库根而非农场目录。
        output_dir = fallback_output_dir("word-extractor", "doc-output")
    output_dir.mkdir(parents=True, exist_ok=True)

    if args.format in ("json", "both"):
        json_path = output_dir / f"{base_name}.json"
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        print(f"  JSON: {json_path}")

    if args.format in ("markdown", "both"):
        md = generate_markdown(data)
        md_path = output_dir / f"{base_name}.md"
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md)
        print(f"  Markdown: {md_path}")

    # Print summary
    stats = data["statistics"]
    print(f"\nSummary:")
    print(f"  Paragraphs: {stats['paragraph_count']}")
    print(f"  Headings: {stats['heading_count']}")
    print(f"  Tables: {stats['table_count']}")
    print(f"  Images: {stats['image_count']}")
    print(f"  Comments: {stats['comment_count']}")
    if data["sections"]["title"]:
        print(f"  Title: {data['sections']['title']}")


if __name__ == "__main__":
    main()
