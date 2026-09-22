#!/usr/bin/env python3
"""Tests for word-extractor extract_docx.py module.

Covers:
  1. Pure-text .docx — paragraph extraction
  2. Table .docx — table markdown output, including _handle_merged_cells
  3. Image .docx — image metadata extraction
  4. Comment .docx — comment extraction

Run:
    cd ~/projects/dc-skills && uv run pytest word-extractor/scripts/test_extract_docx.py -v
"""

from __future__ import annotations

import re
import struct
import sys
import zlib
import zipfile
from pathlib import Path

from docx import Document

# Ensure the word-extractor scripts dir is on sys.path
_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

import extract_docx as ed  # noqa: E402


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_1x1_png() -> bytes:
    """Create a minimal 1x1 red PNG in memory (valid PNG binary)."""
    # Build a minimal valid PNG: 1x1 pixel, RGB, red (255,0,0)
    def chunk(chunk_type: bytes, data: bytes) -> bytes:
        crc_input = chunk_type + data
        crc = struct.pack(">I", zlib.crc32(crc_input) & 0xFFFFFFFF)
        return struct.pack(">I", len(data)) + chunk_type + data + crc

    # IHDR: 1x1, 8-bit RGB
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    # IDAT: filtered scanline (filter=0, then RGB=255,0,0)
    raw_scanline = b"\x00\xff\x00\x00"  # filter byte 0 + RGB red
    idat = zlib.compress(raw_scanline)
    signature = b"\x89PNG\r\n\x1a\n"
    return signature + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


def _is_text_entry(filename: str) -> bool:
    return filename.endswith(".xml") or filename.endswith(".rels")


def _inject_comment(docx_path: Path, comment_id: int = 0,
                    author: str = "Tester", text: str = "Test comment",
                    initals: str = "TT") -> None:
    """Inject a comment into a saved .docx via raw ZIP manipulation.

    Adds a comments.xml part, updates [Content_Types].xml and
    document.xml.rels so python-docx can find it on re-open.
    Then injects commentRangeStart/End markers into the first paragraph
    of document.xml.
    """
    comments_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">\n'
        f'  <w:comment w:id="{comment_id}" w:author="{author}" '
        f'w:date="2024-01-01T00:00:00Z" w:initials="{initals}">\n'
        '    <w:p><w:r><w:t>' + text + '</w:t></w:r></w:p>\n'
        '  </w:comment>\n'
        '</w:comments>'
    )

    range_start_tag = f'<w:commentRangeStart w:id="{comment_id}"/>'
    range_end_tag = f'<w:commentRangeEnd w:id="{comment_id}"/>'

    tmp_path = Path(str(docx_path) + ".tmp")

    rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
    content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"

    with zipfile.ZipFile(str(docx_path), "r") as zin:
        with zipfile.ZipFile(str(tmp_path), "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if not _is_text_entry(item.filename):
                    zout.writestr(item, data)
                    continue
                data_str = data.decode("utf-8")
                if item.filename == "[Content_Types].xml":
                    data_str = data_str.replace(
                        "</Types>",
                        f'<Override PartName="/word/comments.xml" '
                        f'ContentType="{content_type}"/></Types>')
                elif item.filename == "word/_rels/document.xml.rels":
                    data_str = data_str.replace(
                        "</Relationships>",
                        f'<Relationship Id="rIdCom99" Type="{rel_type}" '
                        f'Target="comments.xml"/></Relationships>')
                elif item.filename == "word/document.xml":
                    data_str = _inject_range_markers(
                        data_str, range_start_tag, range_end_tag)
                zout.writestr(item, data_str.encode("utf-8"))
            zout.writestr("word/comments.xml", comments_xml.encode("utf-8"))

    tmp_path.rename(docx_path)


def _inject_range_markers(doc_xml: str, start_tag: str, end_tag: str) -> str:
    """Inject commentRangeStart/End into the first <w:p> of doc XML.  Complex
    manual tag matching here because we work on raw strings, not a parsed DOM."""
    para_pattern = re.compile(r'(<w:p\b[^>]*>)')
    match = para_pattern.search(doc_xml)
    if not match:
        return doc_xml

    para_start = match.end()

    # Manual depth counting to find matching </w:p>
    depth = 1
    pos = para_start
    para_end = -1
    while depth > 0 and pos < len(doc_xml):
        next_open = doc_xml.find("<w:p", pos + 1)
        next_close = doc_xml.find("</w:p>", pos)
        if next_close == -1:
            break
        if 0 < next_open < next_close:
            depth += 1
            pos = next_open + 1
        else:
            depth -= 1
            if depth == 0:
                para_end = next_close
                break
            pos = next_close + 1

    if para_end < 0:
        return doc_xml

    para_content = doc_xml[para_start:para_end]
    r_positions = list(re.finditer(r'<w:r\b[^>]*>', para_content))
    if r_positions:
        last_r_end = 0
        for rp in r_positions:
            r_depth = 1
            p = rp.end()
            while r_depth > 0 and p < len(para_content):
                no = para_content.find("<w:r", p)
                nc = para_content.find("</w:r>", p)
                if nc == -1:
                    break
                if 0 < no < nc:
                    r_depth += 1
                    p = no + 1
                else:
                    r_depth -= 1
                    if r_depth == 0:
                        last_r_end = nc + len("</w:r>")
                        break
                    p = nc + 1
        insert_pos = para_start + last_r_end
        return doc_xml[:para_start] + start_tag + para_content[:last_r_end] + end_tag + para_content[last_r_end:] + doc_xml[para_end:]
    else:
        return doc_xml[:para_start] + start_tag + end_tag + doc_xml[para_start:]


# ---------------------------------------------------------------------------
# Test 1: Pure-text .docx — paragraph extraction
# ---------------------------------------------------------------------------

def test_extract_pure_text_paragraphs(tmp_path: Path) -> None:
    """Create a simple .docx with paragraphs; verify extract_all returns them."""
    doc = Document()
    doc.add_paragraph("Hello World")
    doc.add_paragraph("Second paragraph with more text.")
    doc.add_paragraph("")   # empty paragraph
    docx_path = tmp_path / "pure_text.docx"
    doc.save(str(docx_path))

    result = ed.extract_all(str(docx_path))
    stats = result["statistics"]
    paragraphs = result["paragraphs"]

    assert stats["paragraph_count"] >= 3
    assert stats["table_count"] == 0
    assert stats["image_count"] == 0
    assert stats["comment_count"] == 0

    # Verify paragraph text
    texts = [p["text"] for p in paragraphs]
    assert "Hello World" in texts
    assert "Second paragraph with more text." in texts

    # Verify paragraph metadata
    hello_para = next(p for p in paragraphs if p["text"] == "Hello World")
    assert hello_para["index"] is not None
    assert hello_para["style"] is not None  # "Normal" or similar


# ---------------------------------------------------------------------------
# Test 2: Table .docx — _handle_merged_cells output
# ---------------------------------------------------------------------------

def test_extract_table_with_merged_cells(tmp_path: Path) -> None:
    """Create a .docx with a table that has horizontally merged cells;
    verify _handle_merged_cells runs and places the merged cell content."""
    doc = Document()
    table = doc.add_table(rows=3, cols=3, style="Table Grid")

    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(0, 2).text = "C1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "MERGED"
    table.cell(2, 0).text = "A3"
    table.cell(2, 1).text = "B3"
    table.cell(2, 2).text = "C3"

    # Merge B2 + C2 horizontally (row 1, cols 1-2)
    table.cell(1, 1).merge(table.cell(1, 2))

    docx_path = tmp_path / "table.docx"
    doc.save(str(docx_path))

    result = ed.extract_all(str(docx_path))
    tbl = result["tables"][0]

    assert result["statistics"]["table_count"] == 1
    assert tbl["rows"] == 3
    assert len(tbl["cells"]) == 3

    # All rows must have the same column count
    col_counts = {len(row) for row in tbl["cells"]}
    assert len(col_counts) == 1, f"inconsistent columns: {col_counts}"

    # Row 0: three distinct cells
    assert tbl["cells"][0][0] == "A1"
    assert tbl["cells"][0][1] == "B1"
    assert tbl["cells"][0][2] == "C1"

    # Row 1: merged cell text appears (at col 1)
    assert tbl["cells"][1][0] == "A2"
    assert tbl["cells"][1][1] == "MERGED"


def test_handle_merged_cells_no_merges(tmp_path: Path) -> None:
    """_handle_merged_cells on a simple 2x2 table without merges returns
    the same cells as naive extraction."""
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "R0C0"
    table.cell(0, 1).text = "R0C1"
    table.cell(1, 0).text = "R1C0"
    table.cell(1, 1).text = "R1C1"
    docx_path = tmp_path / "no_merge.docx"
    doc.save(str(docx_path))

    result = ed.extract_all(str(docx_path))
    tbl = result["tables"][0]

    assert tbl["columns"] == 2
    assert tbl["rows"] == 2
    assert tbl["cells"][0] == ["R0C0", "R0C1"]
    assert tbl["cells"][1] == ["R1C0", "R1C1"]


# ---------------------------------------------------------------------------
# Test 3: Image .docx — image metadata extraction
# ---------------------------------------------------------------------------

def test_extract_image_metadata(tmp_path: Path) -> None:
    """Create a .docx with an inline image; verify extract_images captures
    filename, paragraph index, and association."""
    # Create a tiny PNG
    png_data = _make_1x1_png()
    png_path = tmp_path / "tiny.png"
    png_path.write_bytes(png_data)

    doc = Document()
    doc.add_paragraph("Before image paragraph.")
    doc.add_picture(str(png_path), width=None, height=None)
    doc.add_paragraph("After image paragraph.")
    docx_path = tmp_path / "image.docx"
    doc.save(str(docx_path))

    result = ed.extract_all(str(docx_path))
    stats = result["statistics"]
    images = result["images"]

    assert stats["image_count"] >= 1, \
        f"Expected at least 1 image, got {stats['image_count']}"
    assert len(images) >= 1

    img = images[0]
    assert "filename" in img
    assert img["filename"].endswith(".png"), \
        f"Expected .png filename, got {img['filename']}"
    assert "paragraph_index" in img
    assert isinstance(img["paragraph_index"], int)
    assert img["paragraph_index"] >= 0
    assert "rId" in img


# ---------------------------------------------------------------------------
# Test 4: Comment .docx — comment extraction
# ---------------------------------------------------------------------------

def test_extract_comments(tmp_path: Path) -> None:
    """Create a .docx, inject a comment via ZIP manipulation, then verify
    extract_comments_data and extract_comment_ranges parse it correctly."""
    doc = Document()
    p = doc.add_paragraph("This paragraph has a comment attached.")
    doc.add_paragraph("Another paragraph without comment.")

    docx_path = tmp_path / "comment.docx"
    doc.save(str(docx_path))

    # Inject a comment with author + text
    _inject_comment(docx_path, comment_id=0, author="Reviewer",
                    text="This looks good.", initals="RV")

    # Now re-open and extract
    doc2 = Document(str(docx_path))
    comments_map = ed.extract_comments_data(doc2)
    comment_paras, comment_texts = ed.extract_comment_ranges(doc2)

    assert "0" in comments_map, \
        f"Expected comment id 0 in map, got keys: {list(comments_map.keys())}"
    c = comments_map["0"]
    assert c["author"] == "Reviewer"
    assert c["text"] == "This looks good."
    assert c["initials"] == "RV"

    # comment_paras should map comment 0 to at least paragraph 0
    assert "0" in comment_paras, \
        f"Expected comment 0 in comment_paras, got keys: {list(comment_paras.keys())}"
    assert 0 in comment_paras["0"]
